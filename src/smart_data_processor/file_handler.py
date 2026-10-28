"""Validated, read-only CSV/XLSX imports with conservative value handling."""

import codecs
import csv
from io import BytesIO, StringIO
from os import PathLike
from pathlib import Path
from typing import BinaryIO
from zipfile import BadZipFile, ZipFile

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.utils.exceptions import InvalidFileException
from xml.etree.ElementTree import ParseError

from .validators import DataValidationError, FileValidationError, validate_headers
from .excel_import import ExcelImportPlan, plan_excel_import, preview_import, region_import_options
from .excel_layout import ExcelFormula, WorksheetLayout, analyze_worksheet

DEFAULT_MAX_SIZE_BYTES = 10 * 1024 * 1024
DEFAULT_MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
_MAX_WORKSHEET_CELLS = 2_000_000


def load_dataset(
    source: str | PathLike[str] | BinaryIO,
    *,
    filename: str | None = None,
    max_size_bytes: int = DEFAULT_MAX_SIZE_BYTES,
    encoding: str | None = None,
    delimiter: str | None = None,
    sheet_name: str | int = 0,
    max_uncompressed_bytes: int = DEFAULT_MAX_UNCOMPRESSED_BYTES,
    cell_range: str | None = None,
    exclude_empty_columns: bool = False,
    generate_missing_headers: bool = False,
    deduplicate_headers: bool = False,
    confirm_header_changes: bool = False,
    header_row: int | None = None,
    infer_header: bool = False,
    remove_empty_rows: bool = False,
    repair_irregular_headers: bool = False,
    confirm_import_changes: bool = False,
    headerless: bool = False,
) -> pd.DataFrame:
    """Load a local path or seekable binary stream without changing the source.

    Streams require ``filename`` (or a ``name`` attribute) and their position is
    restored. CSV fields remain strings; only empty fields become ``pd.NA``.
    UTF-8/BOM and BOM-marked UTF-16/32 are detected; ambiguous legacy encodings
    require an explicit override. Delimiters are sniffed from comma, semicolon,
    tab, and pipe. XLSX preserves formulas and native cell values, without
    evaluating formulas. The first worksheet is selected by default.

    XLSX defaults to strict headers. Explicit range/exclusion/header options can
    be reviewed with inspect_xlsx_import; actual renames additionally require
    confirm_header_changes=True. Selected cell values are never rewritten.
    Header-only datasets are valid; empty files, ragged rows, and invalid headers
    are rejected. Limits apply to input bytes and XLSX expanded
    archive bytes. All failures in input validation raise FileValidationError.
    """
    for label, value in (
        ("max_size_bytes", max_size_bytes),
        ("max_uncompressed_bytes", max_uncompressed_bytes),
    ):
        if type(value) is not int or value <= 0:
            raise FileValidationError(f"{label} must be a positive integer.")

    is_path = isinstance(source, (str, PathLike))
    name = str(source) if is_path else filename or getattr(source, "name", "")
    if not isinstance(name, str) or Path(name).suffix.lower() not in {".csv", ".xlsx"}:
        raise FileValidationError("Provide a filename with a .csv or .xlsx extension.")
    suffix = Path(name).suffix.lower()
    excel_switches = (exclude_empty_columns, generate_missing_headers, deduplicate_headers, confirm_header_changes,
                     infer_header, remove_empty_rows, repair_irregular_headers, confirm_import_changes, headerless)
    if any(type(value) is not bool for value in excel_switches):
        raise FileValidationError("Excel import switches must be boolean.")
    if suffix == ".csv" and (cell_range is not None or header_row is not None or any(excel_switches)):
        raise FileValidationError("Excel range and header options are only supported for XLSX files.")
    if suffix == ".csv" and (type(sheet_name) is not int or sheet_name != 0):
        raise FileValidationError("sheet_name is only supported for XLSX files.")
    if suffix == ".xlsx" and (encoding is not None or delimiter is not None):
        raise FileValidationError("encoding and delimiter are only supported for CSV files.")

    try:
        if is_path:
            with Path(source).open("rb") as handle:
                payload = handle.read(max_size_bytes + 1)
        else:
            position = source.tell()
            try:
                source.seek(0)
                payload = source.read(max_size_bytes + 1)
            finally:
                source.seek(position)
    except (OSError, ValueError, AttributeError, TypeError) as exc:
        raise FileValidationError(
            "Cannot read input; provide a readable file or seekable binary stream."
        ) from exc
    if not isinstance(payload, bytes):
        raise FileValidationError("Input streams must return bytes, not text.")
    if len(payload) > max_size_bytes:
        raise FileValidationError(f"Input exceeds the {max_size_bytes}-byte size limit.")
    if not payload:
        raise FileValidationError("The input file is empty.")

    try:
        if suffix == ".csv":
            return _load_csv(payload, encoding, delimiter)
        return _load_xlsx(
            payload, sheet_name, max_uncompressed_bytes, cell_range=cell_range,
            exclude_empty_columns=exclude_empty_columns, generate_missing_headers=generate_missing_headers,
            deduplicate_headers=deduplicate_headers, confirm_header_changes=confirm_header_changes,
            header_row=header_row, infer_header=infer_header, remove_empty_rows=remove_empty_rows,
            repair_irregular_headers=repair_irregular_headers, confirm_import_changes=confirm_import_changes,
            headerless=headerless,
        )
    except DataValidationError as exc:
        raise FileValidationError(str(exc)) from exc


def _decode_csv(payload: bytes, encoding: str | None) -> str:
    if encoding is None:
        if payload.startswith((codecs.BOM_UTF32_LE, codecs.BOM_UTF32_BE)):
            encoding = "utf-32"
        elif payload.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
            encoding = "utf-16"
        else:
            encoding = "utf-8-sig"
    try:
        text = payload.decode(encoding)
    except (UnicodeError, LookupError, TypeError) as exc:
        raise FileValidationError(
            "Cannot decode CSV. Supply its encoding explicitly, e.g. encoding='cp1252'."
        ) from exc
    if "\x00" in text:
        raise FileValidationError("CSV contains NUL characters; check the file and encoding.")
    return text


def _load_csv(payload: bytes, encoding: str | None, delimiter: str | None) -> pd.DataFrame:
    text = _decode_csv(payload, encoding)
    if delimiter is not None and (
        not isinstance(delimiter, str)
        or len(delimiter) != 1
        or delimiter in '\r\n\x00"'
    ):
        raise FileValidationError("delimiter must be one character other than newline, NUL, or quote.")
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|").delimiter
        except csv.Error:
            # Comma is the documented fallback, including single-column inputs.
            delimiter = ","
    try:
        reader = csv.reader(StringIO(text, newline=""), delimiter=delimiter, strict=True)
        headers = next(reader, [])
        validate_headers(headers)
        rows = []
        for row in reader:
            if len(row) != len(headers):
                raise FileValidationError(
                    f"CSV record ending at line {reader.line_num} has {len(row)} fields; "
                    f"expected {len(headers)}. Check the delimiter or malformed rows."
                )
            rows.append([pd.NA if value == "" else value for value in row])
    except csv.Error as exc:
        raise FileValidationError(f"Malformed CSV: {exc}.") from exc
    return pd.DataFrame(rows, columns=headers, dtype=object)


def list_worksheets(
    payload: bytes,
    *,
    max_size_bytes: int = DEFAULT_MAX_SIZE_BYTES,
    max_uncompressed_bytes: int = DEFAULT_MAX_UNCOMPRESSED_BYTES,
) -> tuple[str, ...]:
    """List sheets in uploaded XLSX bytes without requiring a valid first table.

    Uses the same archive validation and resource cleanup as dataset loading.
    Does not read worksheet rows or modify the uploaded bytes.
    """
    _validate_xlsx_payload(payload, max_size_bytes, max_uncompressed_bytes)
    return _load_xlsx(payload, 0, max_uncompressed_bytes, names_only=True)


def _validate_xlsx_payload(payload: bytes, max_size_bytes: int, max_uncompressed_bytes: int) -> None:
    for value in (max_size_bytes, max_uncompressed_bytes):
        if type(value) is not int or value <= 0:
            raise FileValidationError("Size limits must be positive integers.")
    if not isinstance(payload, bytes) or not payload:
        raise FileValidationError("Provide non-empty XLSX bytes.")
    if len(payload) > max_size_bytes:
        raise FileValidationError(f"Input exceeds the {max_size_bytes}-byte size limit.")


def inspect_xlsx_import(
    payload: bytes, *, sheet_name: str | int = 0,
    max_size_bytes: int = DEFAULT_MAX_SIZE_BYTES,
    max_uncompressed_bytes: int = DEFAULT_MAX_UNCOMPRESSED_BYTES,
    cell_range: str | None = None, exclude_empty_columns: bool = False,
    generate_missing_headers: bool = False, deduplicate_headers: bool = False,
    header_row: int | None = None, infer_header: bool = False,
    remove_empty_rows: bool = False, repair_irregular_headers: bool = False,
    headerless: bool = False,
) -> ExcelImportPlan:
    """Preview exact exclusions/renames and validation errors before import.

    Inspection never implies approval. Use the identical options and explicitly
    confirm_header_changes=True when loading a plan that renames headers.
    """
    _validate_xlsx_payload(payload, max_size_bytes, max_uncompressed_bytes)
    return _load_xlsx(
        payload, sheet_name, max_uncompressed_bytes, inspect_only=True,
        cell_range=cell_range, exclude_empty_columns=exclude_empty_columns,
        generate_missing_headers=generate_missing_headers, deduplicate_headers=deduplicate_headers,
        header_row=header_row, infer_header=infer_header, remove_empty_rows=remove_empty_rows,
        repair_irregular_headers=repair_irregular_headers,
        headerless=headerless,
    )


def inspect_xlsx_layouts(
    payload: bytes, *, max_size_bytes: int = DEFAULT_MAX_SIZE_BYTES,
    max_uncompressed_bytes: int = DEFAULT_MAX_UNCOMPRESSED_BYTES,
) -> tuple[WorksheetLayout, ...]:
    """Inspect empty sheets, exact blank margins, suggested regions and formulas."""
    _validate_xlsx_payload(payload, max_size_bytes, max_uncompressed_bytes)
    return _load_xlsx(payload, 0, max_uncompressed_bytes, layouts_only=True)


def load_xlsx_regions(payload: bytes, *, sheet_name: str | int = 0,
                      ranges: tuple[str, ...], **options) -> dict[str, pd.DataFrame]:
    """Import explicitly selected rectangles separately, atomically; never join them.

    Options are load_dataset keyword arguments shared by every selected region.
    Different header rows can be inferred per region and explicitly confirmed.
    Single-cell regions retain their value with a generated name, requiring
    confirm_header_changes=True. They never consume the cell as a header.
    """
    if not isinstance(ranges, (tuple, list)) or not ranges or any(not isinstance(item, str) for item in ranges):
        raise FileValidationError("Select at least one region range.")
    if len(set(ranges)) != len(ranges):
        raise FileValidationError("Region ranges must be unique.")
    results = {}
    for region in ranges:
        arguments = region_import_options(region, options)
        results[region] = load_dataset(BytesIO(payload), filename="regions.xlsx", sheet_name=sheet_name,
                                      cell_range=region, **arguments)
    return results


def _worksheet_layout(worksheet) -> WorksheetLayout:
    # A tiny archive can declare an enormous, mostly blank worksheet. Bound
    # materialization before iterating; never truncate a selected dataset.
    cells = (worksheet.max_row or 0) * (worksheet.max_column or 0)

    def check_size(cells):
        if cells > _MAX_WORKSHEET_CELLS:
            raise FileValidationError(
                f"Worksheet {worksheet.title!r} spans {cells:,} cells including blanks, "
                f"exceeding the {_MAX_WORKSHEET_CELLS:,}-cell inspection limit. Remove unused rows/columns "
                "in a copy of the original workbook, or provide a smaller workbook. No data was imported."
            )

    check_size(cells)
    formulas, details = [], []

    def value_of(cell):
        value = cell.value
        if cell.data_type == "f":
            if isinstance(value, str):
                formulas.append((cell.coordinate, value))
                details.append((cell.coordinate, "formula", None))
            else:
                value = ExcelFormula(getattr(value, "text", None), type(value).__name__, getattr(value, "ref", None))
                formulas.append((cell.coordinate, value.expression))
                details.append((cell.coordinate, value.kind, value.ref))
        return value

    # Retain values, not a second full collection of openpyxl Cell objects.
    records, scanned_cells = [], 0
    for row in worksheet.iter_rows():
        scanned_cells += len(row)
        check_size(scanned_cells)  # Also bound workbooks with no declared dimensions.
        records.append(tuple(value_of(cell) for cell in row))
    return analyze_worksheet(records, worksheet.title, formulas, details)


def _load_xlsx(
    payload: bytes, sheet_name: str | int, max_uncompressed_bytes: int,
    *, names_only: bool = False, inspect_only: bool = False, layouts_only: bool = False,
    cell_range: str | None = None, exclude_empty_columns: bool = False,
    generate_missing_headers: bool = False, deduplicate_headers: bool = False,
    confirm_header_changes: bool = False,
    header_row: int | None = None, infer_header: bool = False,
    remove_empty_rows: bool = False, repair_irregular_headers: bool = False,
    confirm_import_changes: bool = False,
    headerless: bool = False,
) -> pd.DataFrame | tuple[str, ...] | ExcelImportPlan | tuple[WorksheetLayout, ...]:
    if not isinstance(sheet_name, str) and type(sheet_name) is not int:
        raise FileValidationError("sheet_name must be a worksheet name or zero-based index.")
    workbook = None
    try:
        with ZipFile(BytesIO(payload)) as archive:
            if sum(info.file_size for info in archive.infolist()) > max_uncompressed_bytes:
                raise FileValidationError("XLSX exceeds the uncompressed archive size limit.")
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=False)
        names = workbook.sheetnames
        if names_only:
            return tuple(names)
        if layouts_only:
            return tuple(_worksheet_layout(workbook[name]) for name in names)
        if isinstance(sheet_name, str):
            if sheet_name not in names:
                raise FileValidationError(f"Unknown worksheet {sheet_name!r}. Available: {names}.")
            worksheet = workbook[sheet_name]
        else:
            if not 0 <= sheet_name < len(names):
                raise FileValidationError(f"Worksheet index out of range. Available: {names}.")
            worksheet = workbook[names[sheet_name]]
        layout = _worksheet_layout(worksheet)
        records = layout.records
        plan = plan_excel_import(
            records, worksheet.title, cell_range=cell_range, exclude_empty_columns=exclude_empty_columns,
            generate_missing_headers=generate_missing_headers, deduplicate_headers=deduplicate_headers,
            header_row=header_row, infer_header=infer_header, remove_empty_rows=remove_empty_rows,
            repair_irregular_headers=repair_irregular_headers,
            headerless=headerless,
        )
        if inspect_only:
            return plan
        if plan.errors:
            raise FileValidationError(" ".join(plan.errors))
        if plan.requires_header_confirmation and not confirm_header_changes:
            raise FileValidationError("Review the import plan and set confirm_header_changes=True to approve header renaming.")
        if (header_row is not None or infer_header or plan.removed_blank_rows) and not confirm_import_changes:
            raise FileValidationError("Review header selection and row exclusions; set confirm_import_changes=True to approve.")
        included = [column for column in plan.columns if not column.action.startswith("exclude_")]
        selected_columns = {col.excel_column for col in included}
        selected_rows = {plan.header_row, *plan.retained_rows}
        formulas = tuple((coordinate, expression) for coordinate, expression in layout.formulas
                         if coordinate_from_string(coordinate)[0] in selected_columns
                         and coordinate_from_string(coordinate)[1] in selected_rows)
        if any(expression is None for coordinate, expression in formulas):
            raise FileValidationError("This selection contains an unsupported array/data-table formula. Use the original workbook instead.")
        formula_coordinates = {coordinate for coordinate, expression in formulas}
        # Object dtype avoids coercing integers with blanks into floats, for example.
        data = preview_import(records, plan)
        data.attrs["xlsx_source"] = {
            "sheet_name": worksheet.title, "range": plan.selected_range,
            "header_row": plan.header_row, "row_numbers": plan.retained_rows,
            "column_letters": tuple(column.excel_column for column in included),
            "formulas": formulas,
            "formula_details": tuple(detail for detail in layout.formula_details if detail[0] in formula_coordinates),
        }
        return data
    except FileValidationError:
        raise
    except (BadZipFile, InvalidFileException, ParseError, KeyError, OSError, ValueError, TypeError, EOFError) as exc:
        raise FileValidationError(f"Cannot read XLSX workbook: {exc}.") from exc
    finally:
        if workbook is not None:
            workbook.close()
