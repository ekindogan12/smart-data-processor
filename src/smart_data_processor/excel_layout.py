"""Conservative worksheet layout discovery and coordinate-preserving previews."""

from dataclasses import dataclass

import pandas as pd
from openpyxl.utils.cell import get_column_letter, range_boundaries


@dataclass(frozen=True)
class ExcelFormula:
    """A non-scalar Excel formula preserved without converting the stored cell.

    Ordinary formulas retain the legacy string API. Array/data-table formulas
    need a typed value so their kind and target range are not silently erased.
    ``str`` is only a display representation; exporters require acknowledgment.
    """

    expression: str | None
    kind: str
    ref: str | None

    def __str__(self) -> str:
        return self.expression if self.expression is not None else f"[{self.kind}: original workbook required]"


def spans(numbers, *, columns=False) -> tuple[str, ...]:
    """Compact sorted one-based coordinates into exact inclusive ranges."""
    groups = []
    for number in sorted(set(numbers)):
        if groups and number == groups[-1][-1] + 1:
            groups[-1].append(number)
        else:
            groups.append([number])
    label = get_column_letter if columns else str
    return tuple(f"{label(group[0])}:{label(group[-1])}" for group in groups)


@dataclass(frozen=True)
class WorksheetLayout:
    sheet_name: str
    rows: int
    columns: int
    empty: bool
    blank_rows: tuple[str, ...]
    blank_columns: tuple[str, ...]
    leading_rows: tuple[str, ...]
    trailing_rows: tuple[str, ...]
    leading_columns: tuple[str, ...]
    trailing_columns: tuple[str, ...]
    regions: tuple[str, ...]
    formulas: tuple[tuple[str, str | None], ...]
    records: tuple[tuple, ...]
    formula_details: tuple[tuple[str, str, str | None], ...] = ()

    def preview(self, cell_range: str | None = None) -> pd.DataFrame:
        """Full value preview with actual Excel row/column labels; never truncated."""
        left, top, right, bottom = (
            range_boundaries(cell_range) if cell_range else (1, 1, self.columns, self.rows)
        )
        return pd.DataFrame(
            [row[left - 1:right] for row in self.records[top - 1:bottom]],
            index=range(top, bottom + 1),
            columns=[get_column_letter(col) for col in range(left, right + 1)], dtype=object,
        )


def analyze_worksheet(records, sheet_name, formulas=(), formula_details=()) -> WorksheetLayout:
    """Suggest rectangles around spatially connected cells, not semantic tables.

    Eight-neighbor components recognize detached formulas even when there is no
    fully blank separator across the whole sheet. Overlapping bounding rectangles
    are merged to avoid counting any cell twice. Sparse logical tables can still
    produce multiple suggestions; users review them. Formatting-only cells are blank.
    """
    records = tuple(tuple(row) for row in records)
    height = len(records)
    width = len(records[0]) if records else 0
    occupied = {(r, c) for r, row in enumerate(records, 1)
                for c, value in enumerate(row, 1) if value is not None}
    used_rows, used_cols = {r for r, c in occupied}, {c for r, c in occupied}
    blank_rows = set(range(1, height + 1)) - used_rows
    blank_cols = set(range(1, width + 1)) - used_cols
    rectangles = []
    remaining = occupied.copy()
    while remaining:
        start = remaining.pop()
        pending, points = [start], [start]
        while pending:
            row, col = pending.pop()
            for dr, dc in ((-1, -1), (-1, 0), (-1, 1), (0, -1), (0, 1), (1, -1), (1, 0), (1, 1)):
                neighbor = (row + dr, col + dc)
                if neighbor in remaining:
                    remaining.remove(neighbor)
                    pending.append(neighbor)
                    points.append(neighbor)
        rows, cols = {r for r, c in points}, {c for r, c in points}
        rectangle = (min(rows), min(cols), max(rows), max(cols))
        # A sparse component can enclose another component. Merge their boxes,
        # repeating because an enlarged box may overlap an earlier rectangle.
        changed = True
        while changed:
            changed = False
            for other in rectangles[:]:
                if not (rectangle[2] < other[0] or other[2] < rectangle[0]
                        or rectangle[3] < other[1] or other[3] < rectangle[1]):
                    rectangles.remove(other)
                    rectangle = (min(rectangle[0], other[0]), min(rectangle[1], other[1]),
                                 max(rectangle[2], other[2]), max(rectangle[3], other[3]))
                    changed = True
        rectangles.append(rectangle)
    regions = tuple(f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"
                    for r1, c1, r2, c2 in sorted(rectangles))
    return WorksheetLayout(
        sheet_name, height, width, not occupied, spans(blank_rows), spans(blank_cols, columns=True),
        spans(r for r in blank_rows if not used_rows or r < min(used_rows)),
        spans(r for r in blank_rows if used_rows and r > max(used_rows)),
        spans((c for c in blank_cols if not used_cols or c < min(used_cols)), columns=True),
        spans((c for c in blank_cols if used_cols and c > max(used_cols)), columns=True),
        regions, tuple(formulas), records, tuple(formula_details),
    )


def infer_header_row(records, left, top, right, bottom) -> int:
    """Suggest the earliest row with the most text labels in the first 20 rows.

    Formula expressions do not count as labels. This is a reviewable heuristic,
    never evidence that earlier rows may be discarded without confirmation.
    """
    candidates = range(top, min(bottom, top + 19) + 1)
    return max(candidates, key=lambda row: sum(
        isinstance(value, str) and bool(value.strip()) and not value.startswith("=")
        for value in records[row - 1][left - 1:right]
    ))
