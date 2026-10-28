"""Auditable XLSX range/header planning, independent of file I/O and UI."""

from dataclasses import dataclass
import re
import pandas as pd

from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries

from .validators import FileValidationError
from .excel_layout import infer_header_row


@dataclass(frozen=True)
class ExcelColumnPlan:
    """One worksheet column, including exclusions and unchanged headers."""

    position: int
    excel_column: str
    original_header: object
    imported_header: str | None
    completely_empty: bool
    has_data: bool
    action: str


@dataclass(frozen=True)
class ExcelImportPlan:
    """Exact selected rectangle and column decisions; no cell values are changed."""

    sheet_name: str
    selected_range: str
    header_row: int
    last_row: int
    worksheet_rows: int
    columns: tuple[ExcelColumnPlan, ...]
    errors: tuple[str, ...]
    retained_rows: tuple[int, ...] = ()
    removed_blank_rows: tuple[int, ...] = ()
    excluded_rows: tuple[int, ...] = ()
    inferred_header: bool = False

    @property
    def row_count(self) -> int:
        return len(self.retained_rows)

    @property
    def requires_header_confirmation(self) -> bool:
        return any(column.action.startswith("rename_") for column in self.columns)

    @property
    def requires_confirmation(self) -> bool:
        return self.requires_header_confirmation or any(
            column.action.startswith("exclude_") for column in self.columns
        ) or self.header_row != 1 or self.last_row != self.worksheet_rows or bool(self.removed_blank_rows) or self.inferred_header


def preview_import(records, plan: ExcelImportPlan) -> pd.DataFrame:
    """Materialize the exact planned values for preview/load, without confirmation."""
    if plan.errors:
        raise FileValidationError(" ".join(plan.errors))
    included = [column for column in plan.columns if not column.action.startswith("exclude_")]
    return pd.DataFrame(
        [tuple(records[row - 1][column.position - 1] for column in included) for row in plan.retained_rows],
        columns=[column.imported_header for column in included], dtype=object,
    )


def region_import_options(region: str, options: dict) -> dict:
    """Single-cell regions are values, not inferred headers; shared UI/API policy."""
    try:
        left, top, right, bottom = range_boundaries(region.strip().upper())
    except (ValueError, TypeError, AttributeError) as exc:
        raise FileValidationError("Invalid region range.") from exc
    singleton = left == right and top == bottom and top is not None
    return {**options, **({"headerless": True, "header_row": None, "infer_header": False} if singleton else {})}


def _bounds(cell_range: str | None, width: int, height: int) -> tuple[int, int, int, int]:
    if cell_range is None:
        return 1, 1, width, height
    if not isinstance(cell_range, str):
        raise FileValidationError("cell_range must be a range such as A:B or L1:M20.")
    match = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]*)?:([A-Z]{1,3})([1-9][0-9]*)?", cell_range.strip().upper())
    if not match or bool(match[2]) != bool(match[4]):
        raise FileValidationError("Use one contiguous column range (A:B) or cell range (L1:M20).")
    left, right = column_index_from_string(match[1]), column_index_from_string(match[3])
    top, bottom = (int(match[2]), int(match[4])) if match[2] else (1, height)
    if not (1 <= left <= right <= min(width, 16_384) and 1 <= top <= bottom <= min(height, 1_048_576)):
        raise FileValidationError("The range is reversed or outside the worksheet's reported bounds.")
    return left, top, right, bottom


def plan_excel_import(
    records: list[tuple], sheet_name: str, *, cell_range: str | None = None,
    exclude_empty_columns: bool = False, generate_missing_headers: bool = False,
    deduplicate_headers: bool = False,
    header_row: int | None = None, infer_header: bool = False,
    remove_empty_rows: bool = False, repair_irregular_headers: bool = False,
    headerless: bool = False,
) -> ExcelImportPlan:
    """Plan import without filtering/mutating records.

    Completely empty means every cell in the WHOLE worksheet column is None,
    including its header. Whitespace, zero, False and formula text are values.
    Missing headers are None or blank strings; other non-string headers remain
    invalid unless irregular-header repair is explicitly requested.
    Generated/suffixed names reserve every existing non-empty name so
    they cannot overwrite or unnecessarily rename a later original column.
    """
    for option in (exclude_empty_columns, generate_missing_headers, deduplicate_headers,
                   infer_header, remove_empty_rows, repair_irregular_headers, headerless):
        if type(option) is not bool:
            raise FileValidationError("Excel import switches must be boolean.")
    width, height = max((len(row) for row in records), default=0), len(records)
    if not width or not height:
        raise FileValidationError("The dataset must contain at least one column.")
    left, top, right, bottom = _bounds(cell_range, width, height)
    if header_row is not None and (type(header_row) is not int or not top <= header_row <= bottom):
        raise FileValidationError("header_row must be an Excel row number within the selected range.")
    if infer_header and header_row is not None:
        raise FileValidationError("Choose an explicit header_row or infer_header, not both.")
    if headerless and (header_row is not None or infer_header):
        raise FileValidationError("Headerless imports cannot also select or infer a header row.")
    top = infer_header_row(records, left, top, right, bottom) if infer_header else (header_row or top)
    headers = (None,) * width if headerless else records[top - 1]
    empty = [all(row[index] is None for row in records) for index in range(width)]
    retained = [index for index in range(left - 1, right) if not (exclude_empty_columns and empty[index])]

    def repaired(value):
        if value is None or isinstance(value, str) and not value.strip():
            return value
        return " ".join(str(value).split()) if repair_irregular_headers else value

    candidates = [repaired(value) for value in headers]
    reserved = {candidates[index] for index in retained if isinstance(candidates[index], str) and candidates[index].strip()}
    used, columns, errors = set(), [], []
    generated_number = 1
    for index, header in enumerate(candidates):
        letter = get_column_letter(index + 1)
        imported, action = None, "keep"
        has_data = any(row[index] is not None for row_number, row in enumerate(records, 1) if headerless or row_number != top)
        if not left - 1 <= index < right:
            action = "exclude_range"
        elif exclude_empty_columns and empty[index]:
            action = "exclude_empty"
        elif header is None or isinstance(header, str) and not header.strip():
            if generate_missing_headers or headerless:
                while f"Unnamed_Column_{generated_number}" in reserved | used:
                    generated_number += 1
                imported = f"Unnamed_Column_{generated_number}"
                generated_number += 1
                action = "rename_missing"
            else:
                action = "invalid"
                errors.append(f"Column names must be non-empty strings: Excel column {letter} has a missing header.")
        elif not isinstance(header, str):
            action = "invalid"
            errors.append(f"Column names must be non-empty strings: Excel column {letter} has a non-text header.")
        elif header in used:
            if deduplicate_headers:
                suffix = 2
                while f"{header}__{suffix}" in reserved | used:
                    suffix += 1
                imported, action = f"{header}__{suffix}", "rename_duplicate"
            else:
                action = "invalid"
                errors.append(f"Column names must be unique: Excel column {letter} repeats {header!r}.")
        else:
            imported = header
            if repair_irregular_headers and (type(headers[index]) is not str or header != headers[index]):
                action = "rename_irregular"
        if imported is not None:
            used.add(imported)
        columns.append(ExcelColumnPlan(index + 1, letter, headers[index], imported, empty[index], has_data, action))
    if not retained:
        errors.append("The dataset must contain at least one column after exclusions.")
    # Only globally empty rows qualify; cells outside the chosen columns count.
    data_start = top if headerless else top + 1
    removed = tuple(row for row in range(data_start, bottom + 1)
                    if remove_empty_rows and all(value is None for value in records[row - 1]))
    removed_set = set(removed)
    retained_rows = tuple(row for row in range(data_start, bottom + 1) if row not in removed_set)
    excluded = tuple(range(1, top)) + removed + tuple(range(bottom + 1, height + 1))
    return ExcelImportPlan(
        sheet_name, f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}",
        0 if headerless else top, bottom, height, tuple(columns), tuple(errors), retained_rows, removed, excluded, infer_header,
    )
