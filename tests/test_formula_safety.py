"""Phase 4: source references must never become executable relocated formulas."""

from copy import deepcopy
from io import BytesIO, StringIO
import csv
import json
from zipfile import ZipFile
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
import pandas as pd
from pandas.testing import assert_frame_equal
import pytest

from smart_data_processor import (
    DataValidationError, FileValidationError, ProcessingPipeline, ProcessingStep,
    export_csv, export_xlsx, inspect_xlsx_layouts, list_worksheets, load_dataset,
)
from smart_data_processor.cleaner import handle_missing, remove_duplicates
from smart_data_processor.dashboard_state import ImportOptions, Workspace
from smart_data_processor.excel_layout import ExcelFormula
from smart_data_processor.transformer import clean_text


def workbook_payload(expression):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    for address, value in {
        "D6": "id", "E6": "quantity", "F6": "quantity", "H6": "result",
        "D7": "001", "E7": 2, "F7": 3, "H7": expression,
        "D9": "001", "E9": 2, "F9": 3, "H9": expression,
        "D10": "002", "E10": 4, "H10": expression,
    }.items():
        sheet[address] = value
    workbook.create_sheet("Other")["A1"] = 99
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def import_table(payload):
    return load_dataset(BytesIO(payload), filename="source.xlsx", cell_range="D6:H10",
                        exclude_empty_columns=True, remove_empty_rows=True,
                        deduplicate_headers=True, confirm_header_changes=True,
                        confirm_import_changes=True)


def assert_no_executable_formulas(content):
    with ZipFile(BytesIO(content)) as archive:
        for name in archive.namelist():
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                root = ElementTree.fromstring(archive.read(name))
                assert not root.findall(".//{*}f")


@pytest.mark.parametrize("expression", [
    "=E7*F7", "=$E$7*F$7+$F7", "=SUM(E7:F10)", "=A1", "=E8+F8",
    "='Other'!A1", "='[external.xlsx]Sheet1'!A1", "=MissingName", "=Table1[Total]",
    '=INDIRECT("E7")', "=#REF!+1", '=IF(E7=2,"Ana","Ben")', "=SEQUENCE(2)",
])
def test_range_filter_and_header_repair_never_activate_source_references(expression):
    payload = workbook_payload(expression)
    data = import_table(payload)
    original = data.copy(deep=True)
    attrs = deepcopy(data.attrs)
    assert list(data.columns) == ["id", "quantity", "quantity__2", "result"]
    assert data.attrs["xlsx_source"]["column_letters"] == ("D", "E", "F", "H")
    assert data.attrs["xlsx_source"]["row_numbers"] == (7, 9, 10)
    assert data["result"].tolist() == [expression] * 3
    result = ProcessingPipeline([
        ProcessingStep("remove_duplicates"),
        ProcessingStep("handle_missing", {"columns": ["quantity__2"], "strategy": "drop"}),
    ]).run(data)
    assert result.data.shape == (1, 4)
    assert result.data.attrs["xlsx_source"]["row_numbers"] == (7,)
    for exporter in (export_csv, export_xlsx):
        with pytest.raises(DataValidationError, match="acknowledge_formula_loss"):
            exporter(result.data)
    artifact = export_xlsx(result.data, acknowledge_formula_loss=True)
    assert_no_executable_formulas(artifact.content)
    book = load_workbook(BytesIO(artifact.content), data_only=False)
    try:
        cell = book.active["D2"]  # Source H7 moved; source E7/F7 are absent here.
        assert cell.value == expression and cell.data_type == "s"
        assert "missing or moved" in cell.comment.text
        assert "not translated or validated" in cell.comment.text
        assert book.active["A2"].value == "001"
        assert book.sheetnames == ["Processed Data"]
    finally:
        book.close()
    csv_export = export_csv(result.data, acknowledge_formula_loss=True)
    rows = list(csv.reader(StringIO(csv_export.content.decode("utf-8"))))
    assert rows[1] == ["001", "2", "3", "'" + expression]
    assert "original workbook" in artifact.warnings[0].message
    assert_frame_equal(data, original)
    assert data.attrs == attrs
    assert import_table(payload).equals(original)


def test_actual_eko_formula_references_are_literal_in_new_coordinates(sayfa1_workbook):
    data = load_dataset(BytesIO(sayfa1_workbook), filename="sayfa1.xlsx", sheet_name="eko", cell_range="K6:N9")
    assert data.attrs["xlsx_source"]["formulas"][0] == ("N7", "=L7+M7")
    artifact = export_xlsx(data, acknowledge_formula_loss=True)
    assert_no_executable_formulas(artifact.content)
    book = load_workbook(BytesIO(artifact.content))
    try:
        assert book.active.max_row == 4 and book.active.max_column == 4
        assert book.active["D2"].value == "=L7+M7"  # Never pretend this is =B2+C2.
        assert book.active["D2"].data_type == "s"
        assert book.active["D2"].comment
        assert book.active["B2"].value == 32 and book.active["C2"].value == 13
    finally:
        book.close()


@pytest.mark.parametrize("operation", ["trim", "lower", "upper", "title"])
def test_text_cleaning_preserves_formula_definitions_and_quoted_literals(operation):
    expression = '=IF(E7=2," Ana ","Ben")'
    data = import_table(workbook_payload(expression))
    before = data.copy(deep=True)
    result = clean_text(data, columns=["result"], operation=operation)
    assert_frame_equal(result.data, before)
    assert [warning.code for warning in result.warnings] == ["formula_preserved"] * 3
    assert not result.affected_positions and not result.errors
    assert_frame_equal(data, before)


def test_formula_protection_is_conservative_for_copies_but_does_not_change_csv_text_api():
    data = import_table(workbook_payload('="Ana"'))
    data.iloc[0, 0] = '="Ana"'  # Literal copy: exact source expression stays protected.
    assert clean_text(data, columns=["id"], operation="upper").data.iloc[0, 0] == '="Ana"'
    literal = pd.DataFrame({"text": ['="Ana"']})
    assert clean_text(literal, columns=["text"], operation="upper").data.iloc[0, 0] == '="ANA"'


def test_case_insensitive_duplicates_do_not_casefold_formula_literals():
    payload = workbook_payload('="Ana"')
    data = import_table(payload)
    data.loc[data.index[1], "result"] = '="ANA"'
    data.attrs["xlsx_source"]["formulas"] += (("H9", '="ANA"'),)
    filtered = remove_duplicates(data, subset=["result"], case_insensitive_columns=["result"])
    assert filtered["result"].tolist() == ['="Ana"', '="ANA"']
    assert filtered.attrs["xlsx_source"]["row_numbers"] == (7, 9)


def test_standalone_filters_keep_source_coordinates_with_repeated_index_labels():
    data = import_table(workbook_payload("=E7*F7"))
    data.index = [4, 4, 4]
    before = deepcopy(data.attrs)
    deduplicated = remove_duplicates(data, keep="last")
    assert deduplicated.attrs["xlsx_source"]["row_numbers"] == (9, 10)
    filtered = handle_missing(deduplicated, columns=["quantity__2"], strategy="drop").data
    assert filtered.attrs["xlsx_source"]["row_numbers"] == (9,)
    assert data.attrs == before
    assert deduplicated.attrs["xlsx_source"]["row_numbers"] == (9, 10)


def test_typed_array_formula_export_is_annotated_even_without_provenance():
    formula = ExcelFormula("=SUM(A1:A3)", "ArrayFormula", "B1:B3")
    data = pd.DataFrame({"formula": [formula]})
    assert clean_text(data, columns=["formula"], operation="upper").data.iloc[0, 0] == formula
    with pytest.raises(DataValidationError):
        export_xlsx(data)
    artifact = export_xlsx(data, acknowledge_formula_loss=True)
    book = load_workbook(BytesIO(artifact.content))
    try:
        assert book.active["A2"].value == formula.expression
        assert book.active["A2"].data_type == "s" and book.active["A2"].comment
    finally:
        book.close()


def test_export_json_records_acknowledgment_errors_and_reference_limitations(sayfa1_workbook):
    state = Workspace()
    options = ImportOptions(sheet_name="eko", cell_range="K6:N9")
    state.load(sayfa1_workbook, "sayfa1.xlsx", options)
    state.configure([])
    state.make_preview()
    state.apply_preview()
    state.prepare_downloads()
    rejected = json.loads(state.exports["json"].content)
    assert not rejected["formula_loss_acknowledged"]
    assert rejected["exports"]["xlsx"]["status"] == "failed"
    assert "acknowledge_formula_loss" in rejected["exports"]["csv"]["error"]
    state.acknowledge_formula_loss(True)
    state.prepare_downloads()
    report = json.loads(state.exports["json"].content)
    assert report["formula_loss_acknowledged"]
    assert report["exports"]["xlsx"]["status"] == "ready"
    assert "not translated or validated" in report["exports"]["xlsx"]["warnings"][0]["message"]
    assert report["xlsx_source"]["formulas"][0] == ["N7", "=L7+M7"]
    assert state.source_bytes == sayfa1_workbook
    state.configure([ProcessingStep("clean_text", {"columns": ["c"], "operation": "lower"})])
    assert not state.exports and not state.formula_loss_acknowledged and state.applied is None
    with pytest.raises(DataValidationError, match="Confirm"):
        state.prepare_downloads()


def test_small_sparse_workbook_with_huge_dimensions_is_rejected_before_materialization(monkeypatch):
    workbook = Workbook()
    workbook.active["A1"] = "id"
    workbook.active.cell(1_000_001, 3, "detached")
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    payload = buffer.getvalue()
    assert len(payload) < 10_000
    assert list_worksheets(payload) == ("Sheet",)
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", lambda *args, **kwargs: pytest.fail("Must reject before materialization"))
    with pytest.raises(FileValidationError, match="3,000,003 cells.*2,000,000-cell"):
        inspect_xlsx_layouts(payload)
    with pytest.raises(FileValidationError, match="No data was imported"):
        load_dataset(BytesIO(payload), filename="large.xlsx", cell_range="A1:A2")


@pytest.mark.parametrize("limit", [3, 4])
def test_grid_limit_handles_undeclared_dimensions_and_inclusive_boundary(monkeypatch, limit):
    from smart_data_processor import file_handler
    from types import SimpleNamespace

    workbook = Workbook()
    workbook.active.append(["id", "value"])
    workbook.active.append([1, 2])
    # Model a read-only sheet whose producer omitted the optional dimension tag.
    worksheet = SimpleNamespace(max_row=None, max_column=None, title="Data",
                                iter_rows=workbook.active.iter_rows)
    monkeypatch.setattr(file_handler, "_MAX_WORKSHEET_CELLS", limit)
    try:
        if limit == 3:
            with pytest.raises(FileValidationError, match="4 cells.*3-cell"):
                file_handler._worksheet_layout(worksheet)
        else:
            layout = file_handler._worksheet_layout(worksheet)
            assert layout.records == (("id", "value"), (1, 2))
            assert layout.regions == ("A1:B2",)
    finally:
        workbook.close()
