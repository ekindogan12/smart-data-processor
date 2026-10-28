"""Shared in-memory and committed sample fixtures."""

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
import pytest


@pytest.fixture
def sample_directory() -> Path:
    return Path(__file__).resolve().parents[1] / "sample_data"


@pytest.fixture
def workbook_bytes():
    """Build a fresh binary workbook with named worksheets for each test."""
    def build(sheets):
        workbook = Workbook()
        workbook.remove(workbook.active)
        for name, records in sheets.items():
            sheet = workbook.create_sheet(name)
            for record in records:
                sheet.append(record)
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        stream.seek(0)
        return stream
    return build


@pytest.fixture
def two_tables_workbook():
    """Regression layout: A:B and L:M tables, empty C:K and styled empty N."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Measurements"
    for coordinate, value in {
        "A1": "f(Hz.)", "B1": "dB", "L1": "f(Hz.)", "M1": "f(o)",
        "A2": 10, "B2": -3, "A3": 20, "B3": -6, "A4": 40,
        "L2": 1000, "M2": 1.5, "L3": 2000, "M3": 2.5,
    }.items():
        sheet[coordinate] = value
    # Real Excel files often retain a used range because of formatting alone.
    sheet["N4"].number_format = "0.00"
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


@pytest.fixture
def sayfa1_like_workbook(sample_directory):
    """Synthetic margin/blank-column supplement to the real workbook fixture."""
    return (sample_directory / "sayfa1_like.xlsx").read_bytes()


@pytest.fixture
def sayfa1_workbook():
    """Unmodified regression copy of the user's actual OneDrive workbook."""
    return (Path(__file__).parent / "fixtures" / "sayfa1.xlsx").read_bytes()
