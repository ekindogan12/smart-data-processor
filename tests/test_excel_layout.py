"""Realistic sparse-workbook regressions and formula-preservation contracts."""

from io import BytesIO

from openpyxl import load_workbook, Workbook
import pandas as pd
from pandas.testing import assert_frame_equal
import pytest

from smart_data_processor import (
    DataValidationError, FileValidationError, ProcessingPipeline, ProcessingStep,
    export_csv, export_xlsx, inspect_xlsx_import, inspect_xlsx_layouts,
    load_dataset, load_xlsx_regions,
)
from smart_data_processor.excel_import import preview_import
from smart_data_processor.excel_layout import analyze_worksheet


REPAIR = dict(generate_missing_headers=True, deduplicate_headers=True,
              repair_irregular_headers=True, confirm_header_changes=True)


def test_sayfa_like_layout_has_exact_margins_and_separate_regions(sayfa1_like_workbook):
    empty, layout = inspect_xlsx_layouts(sayfa1_like_workbook)
    assert empty.sheet_name == "Sayfa1" and empty.empty and not empty.regions
    assert layout.sheet_name == "eko" and not layout.empty
    assert (layout.rows, layout.columns) == (12, 16)
    assert layout.leading_rows == ("1:3",)
    assert layout.trailing_rows == ("9:12",)
    assert layout.leading_columns == ("A:C",)
    assert layout.trailing_columns == ("O:P",)
    assert layout.blank_columns == ("A:C", "I:J", "O:P")
    assert layout.regions == ("D4:H8", "K4:N7")
    assert ("H5", "=D5*2") in layout.formulas
    assert len(layout.formulas) == 7
    before = layout.preview()
    assert before.loc[5, "H"] == "=D5*2"
    assert before.loc[5, "F"] == "001"


def test_regions_load_separately_preserving_values_and_coordinates(sayfa1_like_workbook):
    tables = load_xlsx_regions(sayfa1_like_workbook, sheet_name="eko",
                               ranges=("D4:H8", "K4:N7"), **REPAIR)
    assert list(tables) == ["D4:H8", "K4:N7"]
    left, right = tables.values()
    assert left.shape == (4, 5) and right.shape == (3, 4)
    assert left.columns.tolist() == ["f(Hz.)", "f(Hz.)__2", "Unnamed_Column_1", "2026", "dB level"]
    assert left.iloc[0].tolist() == [10, 0, "001", True, "=D5*2"]
    assert right.iloc[0].tolist() == [1000, 1.5, "=K5/1000", "sample"]
    assert left.attrs["xlsx_source"]["row_numbers"] == (5, 6, 7, 8)
    assert left.attrs["xlsx_source"]["column_letters"] == ("D", "E", "F", "G", "H")
    left.iloc[0, 0] = "changed"
    assert right.iloc[0, 0] == 1000


def test_combination_is_explicit_aligned_and_never_joins_tables(sayfa1_like_workbook):
    combined = load_dataset(BytesIO(sayfa1_like_workbook), filename="sayfa1.xlsx", sheet_name="eko",
                            header_row=4, exclude_empty_columns=True, remove_empty_rows=True,
                            confirm_import_changes=True, **REPAIR)
    assert combined.shape == (4, 9)
    assert combined.iloc[0].tolist() == [10, 0, "001", True, "=D5*2", 1000, 1.5, "=K5/1000", "sample"]
    assert combined.iloc[-1, 5:].isna().all()
    assert combined.attrs["xlsx_source"]["row_numbers"] == (5, 6, 7, 8)


def test_header_inference_preview_and_confirmation(sayfa1_like_workbook):
    options = dict(sheet_name="eko", cell_range="D:H", infer_header=True,
                   remove_empty_rows=True, generate_missing_headers=True,
                   deduplicate_headers=True, repair_irregular_headers=True)
    plan = inspect_xlsx_import(sayfa1_like_workbook, **options)
    assert plan.header_row == 4 and plan.inferred_header
    assert plan.excluded_rows == (1, 2, 3, 9, 10, 11, 12)
    assert plan.removed_blank_rows == (9, 10, 11, 12)
    assert plan.requires_confirmation and plan.requires_header_confirmation
    with pytest.raises(FileValidationError, match="confirm_import_changes"):
        load_dataset(BytesIO(sayfa1_like_workbook), filename="source.xlsx", confirm_header_changes=True, **options)
    loaded = load_dataset(BytesIO(sayfa1_like_workbook), filename="source.xlsx",
                          confirm_header_changes=True, confirm_import_changes=True, **options)
    layout = inspect_xlsx_layouts(sayfa1_like_workbook)[1]
    assert_frame_equal(loaded, preview_import(layout.records, plan))


@pytest.mark.parametrize("option", [{"header_row": 4}, {"infer_header": True}])
def test_header_choices_cannot_silently_discard_earlier_rows(sayfa1_like_workbook, option):
    with pytest.raises(FileValidationError, match="confirm_import_changes"):
        load_dataset(BytesIO(sayfa1_like_workbook), filename="s.xlsx", sheet_name="eko", cell_range="D:H", **REPAIR, **option)


def test_blank_row_removal_keeps_zero_false_whitespace_and_outside_data(workbook_bytes):
    payload = workbook_bytes({"Sheet": [
        ["id", "outside"], [None, None], [0, None], [False, None], [" ", None],
        ["=1+1", None], [None, "data"], [None, None],
    ]}).getvalue()
    plan = inspect_xlsx_import(payload, cell_range="A:B", remove_empty_rows=True)
    assert plan.removed_blank_rows == (2, 8)
    with pytest.raises(FileValidationError, match="confirm_import_changes"):
        load_dataset(BytesIO(payload), filename="s.xlsx", remove_empty_rows=True)
    result = load_dataset(BytesIO(payload), filename="s.xlsx", remove_empty_rows=True, confirm_import_changes=True)
    assert result.iloc[:, 0].tolist() == [0, False, " ", "=1+1", None]
    plan = inspect_xlsx_import(payload, cell_range="A:A", remove_empty_rows=True)
    assert 7 in plan.retained_rows


def test_irregular_header_repair_reserves_collisions(workbook_bytes):
    payload = workbook_bytes({"Sheet": [[" a ", "a", "a__2", 42, "42", None], [1, 2, 3, 4, 5, 6]]}).getvalue()
    data = load_dataset(BytesIO(payload), filename="s.xlsx", **REPAIR)
    assert data.columns.tolist() == ["a", "a__3", "a__2", "42", "42__2", "Unnamed_Column_1"]
    assert data.iloc[0].tolist() == [1, 2, 3, 4, 5, 6]


@pytest.mark.parametrize("exporter", [export_csv, export_xlsx])
def test_imported_formulas_require_export_acknowledgment(sayfa1_like_workbook, exporter):
    data = load_xlsx_regions(sayfa1_like_workbook, sheet_name="eko", ranges=("D4:H8",), **REPAIR)["D4:H8"]
    original = data.copy(deep=True)
    with pytest.raises(DataValidationError, match="acknowledge_formula_loss"):
        exporter(data)
    artifact = exporter(data, acknowledge_formula_loss=True)
    assert any(w.code == "formula_functionality_loss" for w in artifact.warnings)
    assert_frame_equal(data, original)
    if exporter is export_xlsx:
        workbook = load_workbook(BytesIO(artifact.content), data_only=False)
        assert workbook.active["E2"].value == "=D5*2"
        assert workbook.active["E2"].data_type == "s"
        workbook.close()
    else:
        assert b"'=D5*2" in artifact.content


def test_formula_identity_distinguishes_literal_equals_text():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["literal", "actual"])
    worksheet["A2"] = "=1+1"
    worksheet["A2"].data_type = "s"
    worksheet["B2"] = "=1+1"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    data = load_dataset(BytesIO(stream.getvalue()), filename="f.xlsx")
    assert data.attrs["xlsx_source"]["formulas"] == (("B2", "=1+1"),)


def test_formula_and_coordinate_provenance_survives_pipeline_filtering(sayfa1_like_workbook):
    data = load_xlsx_regions(sayfa1_like_workbook, sheet_name="eko", ranges=("D4:H8",), **REPAIR)["D4:H8"]
    result = ProcessingPipeline([
        ProcessingStep("handle_missing", {"columns": ["Unnamed_Column_1"], "strategy": "drop"}),
        ProcessingStep("normalize_column_names"),
    ]).run(data)
    assert result.data.attrs["xlsx_source"]["row_numbers"] == (5, 6, 8)
    assert data.attrs["xlsx_source"]["row_numbers"] == (5, 6, 7, 8)
    with pytest.raises(DataValidationError, match="formula"):
        export_xlsx(result.data)


def test_recursive_regions_cover_staggered_and_vertically_separated_content():
    records = [
        ("a", "b", None, None), (1, 2, None, None), (None, None, None, None),
        ("c", "d", None, "e"), (3, 4, None, 5),
    ]
    layout = analyze_worksheet(records, "Sheet")
    assert layout.regions == ("A1:B2", "A4:B5", "D4:D5")
    from openpyxl.utils.cell import range_boundaries
    covered = set()
    for region in layout.regions:
        left, top, right, bottom = range_boundaries(region)
        cells = {(r, c) for r in range(top, bottom + 1) for c in range(left, right + 1)}
        assert not cells & covered
        covered |= cells
    assert all((r, c) in covered for r, row in enumerate(records, 1)
               for c, value in enumerate(row, 1) if value is not None)


@pytest.mark.parametrize("options", [
    {"header_row": True}, {"header_row": 0}, {"header_row": 100},
    {"header_row": 4, "infer_header": True}, {"remove_empty_rows": "yes"},
    {"repair_irregular_headers": 1}, {"infer_header": 1},
])
def test_new_options_are_validated(sayfa1_like_workbook, options):
    with pytest.raises(FileValidationError):
        inspect_xlsx_import(sayfa1_like_workbook, sheet_name="eko", **options)


@pytest.mark.parametrize("options", [
    {"header_row": 1}, {"infer_header": True}, {"remove_empty_rows": True},
    {"repair_irregular_headers": True}, {"confirm_import_changes": True},
])
def test_new_excel_options_rejected_for_csv(options):
    with pytest.raises(FileValidationError, match="only supported for XLSX"):
        load_dataset(BytesIO(b"a\n1\n"), filename="s.csv", **options)


def test_separate_regions_fail_atomically_and_source_bytes_unchanged(sayfa1_like_workbook, tmp_path):
    path = tmp_path / "sayfa1_like.xlsx"
    path.write_bytes(sayfa1_like_workbook)
    with pytest.raises(FileValidationError):
        load_xlsx_regions(sayfa1_like_workbook, sheet_name="eko", ranges=("K4:N7", "A1:Z2"), **REPAIR)
    load_dataset(path, sheet_name="eko", cell_range="D4:H8", **REPAIR)
    assert path.read_bytes() == sayfa1_like_workbook


@pytest.mark.parametrize("ranges", [(), ("A1:A2", "A1:A2"), (1,), "A1:A2"])
def test_invalid_region_collections_are_rejected(sayfa1_like_workbook, ranges):
    with pytest.raises(FileValidationError):
        load_xlsx_regions(sayfa1_like_workbook, sheet_name="eko", ranges=ranges)


def test_real_sayfa1_layout_detects_tables_and_detached_formula_regions(sayfa1_workbook):
    layout, empty = inspect_xlsx_layouts(sayfa1_workbook)
    assert layout.sheet_name == "eko"
    assert empty.sheet_name == "Sayfa2" and empty.empty
    assert layout.leading_rows == ("1:5",) and layout.leading_columns == ("A:C",)
    assert not layout.trailing_rows and not layout.trailing_columns
    assert layout.regions == ("D6:H12", "K6:N9", "I8:I8", "P8:P8", "O12:O12", "K13:K13", "O14:O14", "J15:J15")
    assert len(layout.formulas) == 9
    assert ("F8", "ArrayFormula", "F8") in layout.formula_details


def test_real_sayfa1_array_formula_stays_typed_and_exports_require_consent(sayfa1_workbook):
    from smart_data_processor.excel_layout import ExcelFormula
    data = load_dataset(BytesIO(sayfa1_workbook), filename="sayfa1.xlsx", sheet_name="eko",
                        cell_range="D6:H12", **REPAIR)
    formula = data.iloc[1, 2]
    assert isinstance(formula, ExcelFormula)  # Not a string or cached result.
    assert formula.expression == "=F11(D6:D9,7)"
    assert formula.kind == "ArrayFormula" and formula.ref == "F8"
    original = data.copy(deep=True)
    result = ProcessingPipeline([ProcessingStep("remove_duplicates")]).run(data)
    assert result.data.iloc[1, 2] == formula
    with pytest.raises(DataValidationError, match="acknowledge_formula_loss"):
        export_xlsx(result.data)
    artifact = export_xlsx(result.data, acknowledge_formula_loss=True)
    workbook = load_workbook(BytesIO(artifact.content), data_only=False)
    assert workbook.active["C3"].value == formula.expression
    assert workbook.active["C3"].data_type == "s"
    workbook.close()
    assert_frame_equal(data, original)


def test_real_sayfa1_all_regions_preserve_every_original_value(sayfa1_workbook):
    layout = inspect_xlsx_layouts(sayfa1_workbook)[0]
    tables = load_xlsx_regions(sayfa1_workbook, sheet_name="eko", ranges=layout.regions, **REPAIR)
    assert len(tables) == 8
    assert tables["I8:I8"].iloc[0, 0] == "=D9+D8"
    assert tables["K13:K13"].iloc[0, 0] == "istanbul"
    assert tables["I8:I8"].attrs["xlsx_source"]["header_row"] == 0
    assert tables["I8:I8"].attrs["xlsx_source"]["row_numbers"] == (8,)
    from openpyxl.utils.cell import column_index_from_string
    covered = set()
    for data in tables.values():
        source = data.attrs["xlsx_source"]
        for column, letter in enumerate(source["column_letters"]):
            excel_column = column_index_from_string(letter)
            if source["header_row"]:
                covered.add((source["header_row"], excel_column))
            for position, excel_row in enumerate(source["row_numbers"]):
                covered.add((excel_row, excel_column))
                expected = layout.records[excel_row - 1][excel_column - 1]
                assert data.iloc[position, column] == expected or (pd.isna(data.iloc[position, column]) and expected is None)
    assert all((r, c) in covered for r, row in enumerate(layout.records, 1)
               for c, value in enumerate(row, 1) if value is not None)


def test_explicit_headerless_import_retains_first_row(workbook_bytes):
    payload = workbook_bytes({"Sheet": [[1, 2], [3, "=A1"]]}).getvalue()
    plan = inspect_xlsx_import(payload, headerless=True)
    assert plan.header_row == 0 and plan.retained_rows == (1, 2)
    with pytest.raises(FileValidationError, match="confirm_header_changes"):
        load_dataset(BytesIO(payload), filename="s.xlsx", headerless=True)
    data = load_dataset(BytesIO(payload), filename="s.xlsx", headerless=True, confirm_header_changes=True)
    assert data.iloc[0].tolist() == [1, 2]
    assert data.attrs["xlsx_source"]["formulas"] == (("B2", "=A1"),)


def test_typed_formula_cannot_bypass_acknowledgment_when_metadata_is_removed():
    from smart_data_processor.excel_layout import ExcelFormula
    data = pd.DataFrame({"f": [ExcelFormula("=SUM(A1:A2)", "ArrayFormula", "B1:B2")]})
    with pytest.raises(DataValidationError, match="acknowledge_formula_loss"):
        export_csv(data)


def test_unsupported_data_table_formula_offers_original_instead_of_stringifying():
    from openpyxl.worksheet.formula import DataTableFormula
    workbook = Workbook()
    workbook.active.append(["id", "formula"])
    workbook.active["A2"] = 1
    workbook.active["B2"] = DataTableFormula(ref="B2:B3", r1="A2")
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    layout = inspect_xlsx_layouts(stream.getvalue())[0]
    assert layout.formulas == (("B2", None),)
    with pytest.raises(FileValidationError, match="original workbook"):
        load_dataset(BytesIO(stream.getvalue()), filename="s.xlsx")


def test_sayfa1_fixture_hash_is_unmodified(sayfa1_workbook):
    import hashlib
    # Exact hash is documented beside the binary fixture.
    from pathlib import Path
    documented_hash = (Path(__file__).parent / "fixtures" / "sayfa1.sha256").read_text().strip()
    assert hashlib.sha256(sayfa1_workbook).hexdigest() == documented_hash


def test_real_workbook_combined_import_is_explicit_and_preserves_detached_cells(sayfa1_workbook):
    data = load_dataset(BytesIO(sayfa1_workbook), filename="sayfa1.xlsx", sheet_name="eko",
                        header_row=6, exclude_empty_columns=True, confirm_import_changes=True, **REPAIR)
    assert data.shape == (9, 13)
    assert data.attrs["xlsx_source"]["column_letters"] == tuple("DEFGHIJKLMNOP")
    assert len(data.attrs["xlsx_source"]["formulas"]) == 9
    assert data.iloc[8, 6] == "=E8"  # J15 remains at its worksheet-aligned position.
    assert data.iloc[6, 7] == "istanbul"  # K13 is never silently dropped.
    assert data.attrs["xlsx_source"]["row_numbers"] == tuple(range(7, 16))
