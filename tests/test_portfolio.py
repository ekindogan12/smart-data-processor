"""The public demo must reproduce documented values using the real application."""

import json
from pathlib import Path
import subprocess
import sys

from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

ROOT = Path(__file__).resolve().parents[1]
PORTFOLIO = ROOT / "docs" / "portfolio"


def test_upwork_copy_meets_requested_limits_and_matches_documentation():
    copy = json.loads((PORTFOLIO / "upwork.json").read_text(encoding="utf-8"))
    assert 0 < len(copy["title"]) <= 40
    assert 0 < len(copy["description"]) <= 600
    assert len(copy["skills"]) == len(set(copy["skills"])) == 5
    rendered = (PORTFOLIO / "README.md").read_text(encoding="utf-8")
    assert all(copy[field] in rendered for field in ("title", "role", "description"))


def test_demo_script_exports_expected_values_and_preserves_source(tmp_path):
    source = ROOT / "sample_data" / "portfolio_sales_raw.csv"
    before = source.read_bytes()
    subprocess.run([sys.executable, str(ROOT / "scripts" / "run_portfolio_demo.py"),
                    "--output-dir", str(tmp_path)], cwd=ROOT, check=True, capture_output=True)
    assert source.read_bytes() == before
    expected = (ROOT / "sample_data" / "portfolio_sales_expected.csv").read_text(encoding="utf-8")
    assert (tmp_path / "portfolio_sales_cleaned.csv").read_text(encoding="utf-8") == expected
    assert json.loads((tmp_path / "demo_summary.json").read_text()) == json.loads(
        (PORTFOLIO / "demo_expected.json").read_text()
    )
    report = json.loads((tmp_path / "portfolio_sales_report.json").read_text())
    assert report["status"] == "completed" and report["errors"] == []
    assert report["after"]["row_count"] == 10 and len(report["steps"]) == 7
    workbook = load_workbook(tmp_path / "portfolio_sales_cleaned.xlsx", data_only=False)
    try:
        assert (workbook.active.max_row, workbook.active.max_column) == (11, 8)
        assert workbook.active["A2"].value == "000101"
        assert all(cell.data_type != "f" for row in workbook.active for cell in row)
    finally:
        workbook.close()


def test_documented_seven_step_workflow_in_real_streamlit_widgets():
    app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=20).run()
    payload = (ROOT / "sample_data" / "portfolio_sales_raw.csv").read_bytes()
    app.file_uploader[0].upload("portfolio_sales_raw.csv", payload).run()
    app.button(key="load_dataset").click().run()
    workspace = app.session_state["workspace"]
    assert workspace.original.shape == (12, 8)
    assert workspace.original_profile.missing_value_count == 3
    assert workspace.original_profile.duplicate_row_count == 2
    labels = {"clean_text": "Clean text", "remove_duplicates": "Remove duplicates",
              "convert_numeric": "Convert numbers", "handle_missing": "Handle missing values",
              "standardize_dates": "Standardize dates"}
    for definition in json.loads((PORTFOLIO / "demo_steps.json").read_text()):
        operation, params = definition["operation"], definition["parameters"]
        app.selectbox(key="new_operation").select(labels[operation]).run()
        app.button(key="add_operation").click().run()
        key = "step_" + app.session_state["queue"][-1]["id"]
        if "columns" in params:
            app.multiselect(key=f"{key}_columns").set_value(params["columns"]).run()
        if operation == "clean_text":
            app.selectbox(key=f"{key}_text").select(params["operation"]).run()
        elif operation == "handle_missing":
            app.selectbox(key=f"{key}_strategy").select(params["strategy"]).run()
            if params["strategy"] == "constant":
                app.text_input(key=f"{key}_value").set_value(params["value"]).run()
        elif operation == "standardize_dates":
            app.text_area(key=f"{key}_formats").set_value("\n".join(params["input_formats"])).run()
    app.button(key="preview_changes").click().run()
    assert not workspace.preview.report.errors
    assert workspace.preview.data.shape == (10, 8)
    app.checkbox(key="side_by_side").check().run()
    app.checkbox(key=f"confirm_{workspace.revision}").check().run()
    app.button(key="apply_changes").click().run()
    app.button(key="prepare_downloads").click().run()
    assert set(workspace.exports) == {"csv", "xlsx", "json"}
    assert workspace.exports["csv"].content.decode("utf-8") == (
        ROOT / "sample_data" / "portfolio_sales_expected.csv"
    ).read_text(encoding="utf-8")
    report = json.loads(workspace.exports["json"].content)
    assert report["after"]["missing_value_count"] == report["after"]["duplicate_row_count"] == 0
    assert report["exports"]["xlsx"]["warnings"][0]["code"] == "precision_protection"
    assert workspace.original.shape == (12, 8)
    assert not app.exception
