"""Machine-readable reports and formula-safe spreadsheet artifacts."""

import csv
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO, StringIO
import json

from openpyxl import load_workbook
import pandas as pd
from pandas.testing import assert_frame_equal
import pytest

from smart_data_processor import (
    DataValidationError, ProcessingPipeline, ProcessingStep,
    export_csv, export_report, export_xlsx, load_dataset,
)


def test_json_report_has_profiles_steps_counts_diagnostics_and_unicode():
    data = pd.DataFrame({"name": [" Ana García ", "Ben"], "n": ["1", "bad"]})
    result = ProcessingPipeline([
        ProcessingStep("clean_text", {"columns": ["name"]}),
        ProcessingStep("convert_numeric", {"columns": ["n"]}),
    ]).run(data)
    payload = json.loads(export_report(result.report))
    assert payload["schema_version"] == 1
    assert payload["before"]["row_count"] == payload["after"]["row_count"] == 2
    assert payload["steps"][1]["before"] == {"rows": 2, "columns": 2}
    assert payload["affected_rows"] == 1
    assert payload["errors"][0]["code"] == "invalid_numeric"
    assert payload["errors"][0]["row_position"] == 1
    payload["steps"][0]["parameters"]["columns"].append("extra")
    assert result.report.steps[0].parameters["columns"] == ["name"]


def test_report_serializes_decimal_datetime_and_null_configuration():
    data = pd.DataFrame({"n": [None], "date": [None]}, dtype=object)
    result = ProcessingPipeline([
        ProcessingStep("handle_missing", {"columns": ["n"], "strategy": "constant", "value": Decimal("1.25")}),
        ProcessingStep("handle_missing", {"columns": ["date"], "strategy": "constant", "value": datetime(2026, 9, 1)}),
    ]).run(data)
    payload = json.loads(export_report(result.report))
    assert payload["steps"][0]["parameters"]["value"] == "1.25"
    assert payload["steps"][1]["parameters"]["value"] == "2026-09-01T00:00:00"


def test_csv_protects_formulas_headers_and_whitespace_without_mutation():
    strings = ["=1+1", "+SUM(A1:A2)", "-2+3", "@SUM(A1)", "  =1", "\t=1", "\r=1", "\n=1", "\ufeff=1", "＝1+1", "plain", "'already safe"]
    data = pd.DataFrame({"=header": strings, "number": [-12] * len(strings)})
    original = data.copy(deep=True)
    result = export_csv(data)
    rows = list(csv.reader(StringIO(result.content.decode("utf-8"))))
    assert rows[0] == ["'=header", "number"]
    assert [row[0] for row in rows[1:11]] == ["'" + value for value in strings[:10]]
    assert rows[-2][0] == "plain"
    assert rows[-1][0] == "'already safe"
    assert all(row[1] == "-12" for row in rows[1:])
    assert result.protected_cell_count == 11
    assert result.warnings[0].code == "formula_protection"
    assert_frame_equal(data, original)


def test_csv_quotes_embedded_delimiters_newlines_and_quotes():
    data = pd.DataFrame({"text": ['hello,"world"\n=1+1', "Ana García", None], "number": [Decimal("12.50"), 0, -1]}, dtype=object)
    result = export_csv(data, delimiter=";")
    rows = list(csv.reader(StringIO(result.content.decode("utf-8")), delimiter=";"))
    assert rows[1] == ['hello,"world"\n=1+1', "12.50"]
    assert rows[2] == ["Ana García", "0"]
    assert rows[3] == ["", "-1"]
    assert len(rows) == 4


def test_xlsx_formula_text_is_never_formula_and_precision_is_retained():
    data = pd.DataFrame({"=header": ["=HYPERLINK(\"https://example.com\")", "+cmd", "#N/A", None], "number": [9007199254740993, Decimal("0.12345678901234567890123456789"), -2, 1.5]}, dtype=object)
    original = data.copy(deep=True)
    result = export_xlsx(data, sheet_name="Cleaned")
    workbook = load_workbook(BytesIO(result.content), data_only=False)
    sheet = workbook["Cleaned"]
    assert sheet["A1"].value == "=header" and sheet["A1"].data_type == "s"
    assert sheet["A2"].value == data.iloc[0, 0] and sheet["A2"].data_type == "s"
    assert sheet["A4"].value == "#N/A" and sheet["A4"].data_type == "s"
    assert sheet["B2"].value == "9007199254740993" and sheet["B2"].data_type == "s"
    assert sheet["B3"].value == "0.12345678901234567890123456789"
    assert sheet["B4"].value == -2 and sheet["B4"].data_type == "n"
    assert not any(cell.data_type == "f" for row in sheet for cell in row)
    workbook.close()
    assert result.protected_cell_count == 3
    assert {issue.code for issue in result.warnings} == {"formula_protection", "precision_protection"}
    assert_frame_equal(data, original)


@pytest.mark.parametrize("exporter", [export_csv, export_xlsx])
def test_exports_roundtrip_simple_values_and_do_not_include_index(exporter):
    data = pd.DataFrame({"id": ["001", "002"], "name": ["Ana García", "Ben"]}, index=[5, 9])
    extension = "csv" if exporter is export_csv else "xlsx"
    result = exporter(data)
    imported = load_dataset(BytesIO(result.content), filename=f"data.{extension}")
    assert imported.columns.tolist() == ["id", "name"]
    assert imported["id"].tolist() == ["001", "002"]
    assert imported["name"].tolist() == ["Ana García", "Ben"]
    assert not result.warnings


@pytest.mark.parametrize("exporter", [export_csv, export_xlsx])
def test_exports_keep_timezone_information_as_iso_text(exporter):
    value = datetime(2026, 9, 1, 12, tzinfo=timezone.utc)
    result = exporter(pd.DataFrame({"date": [value]}))
    if exporter is export_csv:
        assert "2026-09-01T12:00:00+00:00" in result.content.decode("utf-8")
    else:
        workbook = load_workbook(BytesIO(result.content))
        assert workbook.active["A2"].value == "2026-09-01T12:00:00+00:00"
        workbook.close()


@pytest.mark.parametrize("exporter", [export_csv, export_xlsx])
@pytest.mark.parametrize("value", [float("inf"), Decimal("Infinity"), [1, 2], {"a": 1}, complex(1, 2)])
def test_unsupported_export_values_fail_explicitly(exporter, value):
    with pytest.raises(DataValidationError):
        exporter(pd.DataFrame({"value": [value]}))


@pytest.mark.parametrize(
    "value", ["a" * 32768, "bad\x00text", "bad\x0btext"],
    ids=["oversized_text", "nul_character", "vertical_tab"],
)
def test_xlsx_does_not_silently_truncate_or_remove_illegal_text(value):
    with pytest.raises(DataValidationError, match="text limits|illegal controls"):
        export_xlsx(pd.DataFrame({"value": [value]}))


@pytest.mark.parametrize("name", ["", " ", "x" * 32, "bad/name", "[name]", "'quoted", "bad\x00name"])
def test_invalid_sheet_names(name):
    with pytest.raises(DataValidationError):
        export_xlsx(pd.DataFrame({"value": [1]}), sheet_name=name)


@pytest.mark.parametrize("delimiter", ["", "::", "\n", '"'])
def test_invalid_csv_delimiters(delimiter):
    with pytest.raises(DataValidationError):
        export_csv(pd.DataFrame({"value": [1]}), delimiter=delimiter)


@pytest.mark.parametrize("exporter", [export_csv, export_xlsx])
def test_empty_table_export(exporter):
    data = pd.DataFrame({"id": pd.Series(dtype=object)})
    extension = "csv" if exporter is export_csv else "xlsx"
    result = exporter(data)
    imported = load_dataset(BytesIO(result.content), filename=f"empty.{extension}")
    assert imported.shape == (0, 1)


def test_xlsx_column_limit():
    with pytest.raises(DataValidationError, match="limits"):
        export_xlsx(pd.DataFrame(columns=[f"c{i}" for i in range(16385)]))


def test_export_requires_processing_report():
    with pytest.raises(DataValidationError, match="ProcessingReport"):
        export_report({})


@pytest.mark.parametrize("value", ["bad\ud800text", "bad\ufffetext"])
def test_xlsx_rejects_invalid_xml_unicode(value):
    with pytest.raises(DataValidationError, match="illegal controls"):
        export_xlsx(pd.DataFrame({"value": [value]}, dtype=object))


def test_csv_rejects_invalid_unicode():
    with pytest.raises(DataValidationError, match="invalid Unicode"):
        export_csv(pd.DataFrame({"value": ["bad\ud800text"]}, dtype=object))


def test_invalid_nonfinite_configuration_still_has_a_json_error_report():
    result = ProcessingPipeline([
        ProcessingStep("handle_missing", {"columns": ["n"], "strategy": "constant", "value": float("inf")}),
    ]).run(pd.DataFrame({"n": [None]}, dtype=object))
    payload = json.loads(export_report(result.report))
    assert payload["errors"][0]["code"] == "invalid_operation"
    assert payload["steps"][0]["parameters"]["value"] == "inf"


@pytest.mark.parametrize("exporter", [export_csv, export_xlsx])
def test_large_exact_integer_export_does_not_overflow(exporter):
    value = 10 ** 500
    result = exporter(pd.DataFrame({"n": [value]}, dtype=object))
    if exporter is export_csv:
        assert str(value) in result.content.decode("utf-8")
    else:
        workbook = load_workbook(BytesIO(result.content))
        assert workbook.active["A2"].value == str(value)
        workbook.close()
