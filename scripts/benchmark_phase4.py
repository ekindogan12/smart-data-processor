"""Reproducible local timings; no network, dependencies or generated input files.

Run: python scripts/benchmark_phase4.py --output docs/phase4_benchmark.json
Dataset creation is excluded from timings. Timings are observations, not CI gates.
"""

import argparse
from io import BytesIO
import json
from pathlib import Path
import platform
from time import perf_counter
from importlib.metadata import version

from openpyxl import Workbook, load_workbook
import pandas as pd

from smart_data_processor import (
    ProcessingPipeline, ProcessingStep, export_csv, export_xlsx, inspect_xlsx_layouts,
    load_dataset,
)


def measure():
    results = {"environment": {"platform": platform.platform(), "python": platform.python_version(),
                              **{name: version(name) for name in ("pandas", "openpyxl", "streamlit")}},
               "method": "One run per stage; input generation excluded; wall-clock seconds; no memory measurement.",
               "datasets": []}
    frame = pd.DataFrame({
        "id": [f"{i:08d}" for i in range(50_000)],
        "customer": [" Ana ", " Ben "] * 25_000,
        "quantity": [1, 2, 3, 4, 5] * 10_000,
        "price": ["19.95"] * 50_000,
        "date": ["2026-10-28"] * 50_000,
        "city": ["Madrid", "London"] * 25_000,
        "note": [None, "review"] * 25_000,
        "status": ["ready"] * 50_000,
    })
    frame = pd.concat([frame, frame.iloc[:1000]], ignore_index=True)
    csv_payload = frame.to_csv(index=False).encode("utf-8")
    book = Workbook()
    sheet = book.active
    sheet.title = "Sales"
    for column, name in enumerate([*frame.columns[:7], "total"], 4):
        sheet.cell(6, column, name)
    for row in range(7, 10_007):
        for column, value in enumerate([f"{row - 7:08d}", " Ana ", 2, 19.95,
                                       "2026-10-28", "Madrid", None, f"=F{row}*G{row}"], 4):
            sheet.cell(row, column, value)
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    xlsx_payload = buffer.getvalue()
    for name, payload, options in [
        ("sales.csv", csv_payload, {}),
        ("sales.xlsx", xlsx_payload, {"cell_range": "D6:K10006"}),
    ]:
        timings = {}

        def timed(label, operation):
            start = perf_counter()
            result = operation()
            timings[label] = round(perf_counter() - start, 4)
            return result

        if name.endswith("xlsx"):
            layout = timed("worksheet_analysis", lambda: inspect_xlsx_layouts(payload))[0]
            assert layout.regions == ("D6:K10006",) and len(layout.formulas) == 10_000
        data = timed("load", lambda: load_dataset(BytesIO(payload), filename=name, **options))
        original = data.copy(deep=True)
        result = timed("pipeline", lambda: ProcessingPipeline([
            ProcessingStep("clean_text", {"columns": ["customer"], "operation": "trim"}),
            ProcessingStep("remove_duplicates"),
        ]).run(data))
        assert result.report.status == "completed"
        pd.testing.assert_frame_equal(data, original)
        assert len(result.data) == (10_000 if name.endswith("xlsx") else 50_000)
        outputs = {}
        for kind, exporter in [("csv", export_csv), ("xlsx", export_xlsx)]:
            artifact = timed(f"export_{kind}", lambda: exporter(result.data, acknowledge_formula_loss=True))
            outputs[kind] = len(artifact.content)
            if kind == "xlsx":
                exported = load_workbook(BytesIO(artifact.content), read_only=True, data_only=False)
                assert exported.active.max_row == len(result.data) + 1
                assert all(cell.data_type != "f" for row in exported.active for cell in row)
                exported.close()
        results["datasets"].append({"name": name, "input_bytes": len(payload), "input_shape": list(data.shape),
                                    "output_shape": list(result.data.shape), "seconds": timings,
                                    "export_bytes": outputs, "checks": "Passed: dimensions, non-mutation, no executable exported formulas."})
    return results


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    report = json.dumps(measure(), indent=2)
    if args.output:
        args.output.write_text(report + "\n", encoding="utf-8")
    print(report)
