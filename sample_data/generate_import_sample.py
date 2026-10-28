"""Build a synthetic sayfa1-like regression workbook, not the user's original."""

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook


def build_workbook() -> bytes:
    workbook = Workbook()
    workbook.active.title = "Sayfa1"  # Deliberately empty first worksheet.
    sheet = workbook.create_sheet("eko")
    rows = [
        ["f(Hz.)", "f(Hz.)", None, 2026, "  dB\n level  "],
        [10, 0, "001", True, "=D5*2"],
        [20, "mixed", 3.5, False, "=SUM(D5:D6)"],
        [40, -3, None, "NA", "=D7/2"],
        [80, -6, " ", None, "=D8/2"],
    ]
    for row_number, values in enumerate(rows, 4):
        for column_number, value in enumerate(values, 4):
            sheet.cell(row_number, column_number, value)
    for row_number, values in enumerate([
        ["f(Hz.)", "f(o)", None, "  notes  "],
        [1000, 1.5, "=K5/1000", "sample"],
        [2000, "unknown", "=K6/1000", None],
        [3000, 2.5, "=K7/1000", "last"],
    ], 4):
        for column_number, value in enumerate(values, 11):
            sheet.cell(row_number, column_number, value)
    sheet["P12"].number_format = "0.00"  # Blank trailing used rows/columns.
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


if __name__ == "__main__":
    Path(__file__).with_name("sayfa1_like.xlsx").write_bytes(build_workbook())
